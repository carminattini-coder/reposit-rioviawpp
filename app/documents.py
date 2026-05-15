import io
import logging
from pathlib import Path

import pytesseract
from PIL import Image
from docx import Document
from pypdf import PdfReader

logger = logging.getLogger(__name__)


def extract_text(file_path: str | Path) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    elif suffix in (".docx", ".doc"):
        return _extract_docx(path)
    elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".tiff", ".bmp"):
        return _extract_image(path)
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Formato não suportado: {suffix}")


def extract_text_from_bytes(data: bytes, mime_type: str) -> str:
    if "pdf" in mime_type:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif "word" in mime_type or "docx" in mime_type or "officedocument" in mime_type:
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    elif "image" in mime_type:
        img = Image.open(io.BytesIO(data))
        return pytesseract.image_to_string(img, lang="por+eng")
    elif "text" in mime_type:
        return data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Tipo MIME não suportado: {mime_type}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_image(path: Path) -> str:
    img = Image.open(str(path))
    return pytesseract.image_to_string(img, lang="por+eng")


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]

        # Break at sentence boundary when possible
        if end < len(text):
            last_period = max(chunk.rfind("."), chunk.rfind("\n"))
            if last_period > chunk_size // 2:
                end = start + last_period + 1
                chunk = text[start:end]

        chunks.append(chunk.strip())
        start = end - overlap

    return [c for c in chunks if len(c) > 50]
